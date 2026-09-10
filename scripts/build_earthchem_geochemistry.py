#!/usr/bin/env python3
"""Build KHANAN's EarthChem ECL 4498 geochemical sample and observation layers.

The source workbooks are preserved in their original BIFF/XLS form in the ignored
raw cache. LibreOffice performs a mechanical XLS-to-XLSX conversion in a temporary
directory so openpyxl can parse the cells without modifying the source files.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, MutableMapping, Optional, Sequence, Tuple

import h3
import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
RAW = ROOT / "sources" / "raw" / "earthchem_dutt_2026"
GRID = OUT / "india_mining_prospectivity_grid_h3_r6.csv"
ONTOLOGY = OUT / "india_material_ontology_v1.csv"

DATASET_ID = "SRC_EARTHCHEM_ECL_4498_DUTT_2026"
DATASET_URL = "https://ecl.earthchem.org/view.php?id=4498"
DOI = "10.60520/IEDA/114498"
DOI_URL = "https://doi.org/10.60520/IEDA/114498"
LICENSE = "CC-BY-4.0"
LICENSE_URL = "https://spdx.org/licenses/CC-BY-4.0"
TITLE = (
    "Analytical data of Li-rich pegmatites and associated leucogranites of the "
    "Assam-Meghalaya Gneissic Complex, northeast India"
)
AUTHORS = ["Amrita Dutt", "Debajyoti Paul", "Vineet Goswami", "Dwijesh Ray", "Jeetendra Kumar Dash"]
SOURCE_CITATION = (
    "Dutt, A.; Paul, D.; Goswami, V.; Ray, D.; Dash, J.K. (2026). "
    + TITLE
    + ", Version 1.0. Interdisciplinary Earth Data Alliance (IEDA). "
    + DOI_URL
)
ACCESSED_DATE = "2026-09-10"

ZIP_NAME = "all_dataset_4498.zip"
ZIP_SHA1 = "35918ae2fc602798980bb6b3df914966fc1f0eb4"
BULK_NAME = "4498-1_Duttetal_2026_BulkAnalysis.xls"
INSITU_NAME = "4498-2_Duttetal_2026_InsituAnalysis.xls"
SOURCE_SHA1 = {
    BULK_NAME: "118bc636e3cd7a3b42a109ca4c6985b7a4926f2f",
    INSITU_NAME: "30abe32ae1a77e661342e5083abc98446b3918ef",
}
SOURCE_SHA256 = {
    BULK_NAME: "dc29e57f88ed3846140751d12319bb5f2d446f7baf9b853918e94e3ff0decb30",
    INSITU_NAME: "8d49808ac091772966e7f7cac7280a1d0cd1fde0556bdb9f02fa05c72f7e3604",
}

OXIDE_NAMES = {
    "SiO2": "silicon dioxide",
    "TiO2": "titanium dioxide",
    "Al2O3": "aluminium oxide (alumina)",
    "Fe2O3": "iron(III) oxide",
    "Cr2O3": "chromium(III) oxide",
    "FeO": "iron(II) oxide",
    "MnO": "manganese(II) oxide",
    "MgO": "magnesium oxide",
    "CaO": "calcium oxide",
    "K2O": "potassium oxide",
    "Na2O": "sodium oxide",
    "P2O5": "phosphorus pentoxide",
}
SPECIAL_PARAMETER_NAMES = {
    "87Sr/86Sr": "strontium-87/strontium-86 isotope ratio",
    "143Nd/144Nd": "neodymium-143/neodymium-144 isotope ratio",
    "eNd": "epsilon neodymium",
}
MINERAL_CODE_MAP = {
    "Lpd": ("lepidolite", ""),
    "Ms": ("muscovite", "mineral:muscovite"),
    "Grt": ("garnet group", "group:garnet-group"),
    "Kfs": ("K-feldspar", "group:feldspar-group"),
    "Pl": ("plagioclase", "group:feldspar-group"),
    "Chl": ("chlorite group", "group:chlorite-group"),
    "Bt": ("biotite group", "group:biotite-group"),
    "Ep": ("epidote", "mineral:epidote"),
}
KEY_TRACE_PARAMETERS = ["Li", "Rb", "Cs", "Nb", "Ta", "Co", "Ni", "Cu", "Zn", "Zr", "Th", "U"]

SAMPLE_FIELDS = [
    "record_id",
    "sample_name",
    "record_class",
    "latitude",
    "longitude",
    "coordinate_reference_system",
    "coordinate_precision_note",
    "elevation_m_source",
    "source_location_keyword",
    "source_igsn",
    "source_lithology",
    "source_site_id",
    "h3_r6",
    "state_or_ut_2011",
    "district_2011",
    "census_state_code_2011",
    "census_district_code_2011",
    "grid_geology_age",
    "grid_geology_supergroup",
    "grid_geology_group",
    "grid_geology_stratigraphy",
    "co_located_model_top_materials_json",
    "co_located_model_top_material_scores_json",
    "bulk_analysis_rows",
    "in_situ_analysis_spots",
    "observation_rows_total",
    "numeric_observations",
    "below_detection_observations",
    "unresolved_source_na_observations",
    "source_dash_observations",
    "missing_observations",
    "bulk_parameters_json",
    "in_situ_parameters_json",
    "analyzed_mineral_codes_json",
    "analyzed_mineral_names_curated_json",
    "ontology_material_ids_observed_json",
    "chemical_names_observed_json",
    "li_ppm_bulk",
    "rb_ppm_bulk",
    "cs_ppm_bulk",
    "nb_ppm_bulk",
    "ta_ppm_bulk",
    "co_ppm_bulk",
    "ni_ppm_bulk",
    "cu_ppm_bulk",
    "zn_ppm_bulk",
    "zr_ppm_bulk",
    "th_ppm_bulk",
    "u_ppm_bulk",
    "source_sample_classification_note",
    "sampling_design_limitation",
    "deposit_or_mine_status_verified",
    "model_evidence_role",
    "model_use_status",
    "model_linkage_note",
    "source_dataset_id",
    "source_title",
    "source_doi",
    "source_url",
    "source_authors_json",
    "source_license",
    "source_license_url",
    "source_files_json",
    "source_file_sha1_json",
    "source_file_sha256_json",
    "source_citation",
    "source_accessed_date",
    "data_quality_flags_json",
]

OBSERVATION_FIELDS = [
    "observation_id",
    "sample_record_id",
    "sample_name",
    "source_site_id",
    "latitude",
    "longitude",
    "coordinate_reference_system",
    "h3_r6",
    "state_or_ut_2011",
    "district_2011",
    "analysis_scope",
    "analysis_sequence",
    "source_analyzed_material",
    "source_mineral_code",
    "mineral_name_curated",
    "mineral_ontology_material_id",
    "parameter",
    "parameter_class",
    "chemical_name_curated",
    "chemical_name_basis",
    "component_element_material_ids_json",
    "ontology_material_ids_json",
    "source_value_text",
    "numeric_value",
    "measurement_status",
    "measurement_qualifier",
    "unit_source",
    "unit_normalized",
    "source_method_code_data_sheet",
    "source_method_code_metadata_sheet",
    "technique",
    "instrument",
    "laboratory",
    "analytical_accuracy_source",
    "reference_materials_json",
    "detection_limit_reported",
    "detection_limit_value",
    "detection_limit_unit",
    "procedural_blank_value",
    "procedural_blank_unit",
    "operation_conditions_json",
    "sample_preparation_source",
    "chemical_treatment_source",
    "source_workbook",
    "source_sheet",
    "source_row",
    "source_column",
    "source_dataset_id",
    "source_doi",
    "source_url",
    "source_license",
    "source_file_sha1",
    "source_file_sha256",
    "source_citation",
    "source_accessed_date",
    "deposit_or_mine_status_verified",
    "model_evidence_role",
    "model_use_status",
    "data_quality_flags_json",
]


def compact_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def numeric_text(value: Any) -> str:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        return format(value, ".15g")
    return clean_text(value)


def sha(path: Path, algorithm: str) -> str:
    digest = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download_source() -> None:
    """Record research use, download the CC-BY archive, and verify/extract it."""
    RAW.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers.update({"User-Agent": "KHANAN-open-research/1.0 (+https://github.com/ankitjh4/khanan)"})
    feedback = {
        "use": "Research",
        "email": "",
        "submission_id": "4498",
        "dataDOI": DOI,
        "details[page_num]": "1",
        "details[page_count]": "1",
        "details[finished]": "0",
    }
    response = session.post("https://ecl.earthchem.org/admin/classes/DownloadDB.php", data=feedback, timeout=60)
    response.raise_for_status()
    response = session.post(
        "https://ecl.earthchem.org/dl_multi.php",
        data={"id": "4498", "verifySubmit": "yes", "system": "ec", "dlFormat": ZIP_NAME},
        timeout=120,
    )
    response.raise_for_status()
    zip_path = RAW / ZIP_NAME
    zip_path.write_bytes(response.content)
    if sha(zip_path, "sha1") != ZIP_SHA1:
        raise ValueError("EarthChem archive SHA-1 does not match the publisher checksum")
    with zipfile.ZipFile(zip_path) as archive:
        if sorted(archive.namelist()) != sorted(SOURCE_SHA1):
            raise ValueError(f"Unexpected EarthChem archive members: {archive.namelist()}")
        archive.extractall(RAW)


def ensure_source(refresh: bool) -> None:
    required = [RAW / ZIP_NAME, RAW / BULK_NAME, RAW / INSITU_NAME]
    if refresh or any(not path.exists() for path in required):
        download_source()
    checks = {RAW / ZIP_NAME: ZIP_SHA1, **{RAW / name: value for name, value in SOURCE_SHA1.items()}}
    for path, expected in checks.items():
        actual = sha(path, "sha1")
        if actual != expected:
            raise ValueError(f"SHA-1 mismatch for {path.name}: {actual} != {expected}")
    for name, expected in SOURCE_SHA256.items():
        actual = sha(RAW / name, "sha256")
        if actual != expected:
            raise ValueError(f"SHA-256 mismatch for {name}: {actual} != {expected}")


def convert_xls_sources(destination: Path) -> Dict[str, Path]:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice 'soffice' is required to convert the source XLS workbooks")
    result: Dict[str, Path] = {}
    for name in [BULK_NAME, INSITU_NAME]:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(destination), str(RAW / name)],
            check=True,
            capture_output=True,
            text=True,
        )
        converted = destination / (Path(name).stem + ".xlsx")
        if not converted.exists():
            raise FileNotFoundError(converted)
        result[name] = converted
    return result


def load_ontology() -> Tuple[Dict[str, str], Dict[str, str], set]:
    symbol_to_id: Dict[str, str] = {}
    element_names: Dict[str, str] = {}
    ids = set()
    with ONTOLOGY.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            ids.add(row["material_id"])
            if row["entity_type"] != "element":
                continue
            for symbol in json.loads(row["symbols_or_formulae_json"] or "[]"):
                if re.fullmatch(r"[A-Z][a-z]?", symbol):
                    symbol_to_id[symbol] = row["material_id"]
                    element_names[symbol] = row["material_name"].lower()
    return symbol_to_id, element_names, ids


def extract_symbols(parameter: str) -> List[str]:
    if parameter == "87Sr/86Sr":
        return ["Sr"]
    if parameter in {"143Nd/144Nd", "eNd"}:
        return ["Nd"]
    return list(dict.fromkeys(re.findall(r"[A-Z][a-z]?", parameter)))


def parameter_semantics(
    parameter: str, symbol_to_id: Mapping[str, str], element_names: Mapping[str, str]
) -> Tuple[str, str, str, List[str], List[str]]:
    symbols = extract_symbols(parameter)
    component_ids = [symbol_to_id[symbol] for symbol in symbols if symbol in symbol_to_id]
    if parameter in OXIDE_NAMES:
        category = "major_oxide"
        chemical_name = OXIDE_NAMES[parameter]
        basis = "curated_standard_chemical_name"
        ontology_ids = ["compound:alumina"] if parameter == "Al2O3" else component_ids
    elif parameter in SPECIAL_PARAMETER_NAMES:
        category = "isotope_ratio" if "/" in parameter else "calculated_isotope_value"
        chemical_name = SPECIAL_PARAMETER_NAMES[parameter]
        basis = "curated_measurement_name"
        ontology_ids = component_ids
    elif re.fullmatch(r"[A-Z][a-z]?", parameter):
        category = "trace_element" if parameter != "F" else "element"
        chemical_name = element_names.get(parameter, parameter)
        basis = "ontology_canonical_element_name" if parameter in element_names else "source_parameter"
        ontology_ids = component_ids
    else:
        category = "chemical_parameter"
        chemical_name = parameter
        basis = "source_parameter"
        ontology_ids = component_ids
    return category, chemical_name, basis, component_ids, ontology_ids


def parse_samples(workbook) -> Dict[str, Dict[str, Any]]:
    ws = workbook["2 Samples"]
    rows: Dict[str, Dict[str, Any]] = {}
    for source_row in range(8, ws.max_row + 1):
        name = clean_text(ws.cell(source_row, 1).value)
        if not name:
            continue
        rows[name] = {
            "sample_name": name,
            "igsn": clean_text(ws.cell(source_row, 2).value),
            "latitude": float(ws.cell(source_row, 3).value),
            "longitude": float(ws.cell(source_row, 4).value),
            "elevation": float(ws.cell(source_row, 5).value),
            "location_keyword": clean_text(ws.cell(source_row, 6).value),
            "lithology": clean_text(ws.cell(source_row, 7).value),
            "source_row": source_row,
        }
    return rows


def parse_primary_metadata(workbook, sheet_name: str) -> Dict[str, Dict[str, Any]]:
    ws = workbook[sheet_name]
    output: Dict[str, Dict[str, Any]] = {}
    for source_row in range(7, ws.max_row + 1):
        parameter = clean_text(ws.cell(source_row, 2).value)
        if not parameter:
            continue
        item = output.setdefault(
            parameter,
            {
                "metadata_method_code": clean_text(ws.cell(source_row, 1).value),
                "technique": clean_text(ws.cell(source_row, 4).value),
                "instrument": clean_text(ws.cell(source_row, 5).value),
                "laboratory": clean_text(ws.cell(source_row, 6).value),
                "references": [],
                "accuracy": [],
            },
        )
        if not item["technique"]:
            item["technique"] = clean_text(ws.cell(source_row, 4).value)
            item["instrument"] = clean_text(ws.cell(source_row, 5).value)
            item["laboratory"] = clean_text(ws.cell(source_row, 6).value)
        for start in range(10, ws.max_column + 1, 5):
            reference_name = clean_text(ws.cell(source_row, start).value)
            measured_value = ws.cell(source_row, start + 1).value if start + 1 <= ws.max_column else None
            uncertainty = ws.cell(source_row, start + 2).value if start + 2 <= ws.max_column else None
            uncertainty_unit = clean_text(ws.cell(source_row, start + 3).value) if start + 3 <= ws.max_column else ""
            measurements = ws.cell(source_row, start + 4).value if start + 4 <= ws.max_column else None
            if not reference_name:
                continue
            reference = {
                "name": reference_name,
                "measured_value": measured_value if measured_value is not None else "",
                "reported_uncertainty_or_accuracy": clean_text(uncertainty),
                "reported_uncertainty_unit": uncertainty_unit,
                "number_of_measurements": measurements if measurements is not None else "",
            }
            if reference not in item["references"]:
                item["references"].append(reference)
            accuracy_text = clean_text(uncertainty)
            if "accuracy" in accuracy_text.lower() and accuracy_text not in item["accuracy"]:
                item["accuracy"].append(accuracy_text)
    return output


def parse_method_specific(workbook) -> Dict[str, Dict[str, Any]]:
    ws = workbook["5 Method-specific Metadata"]
    output: Dict[str, Dict[str, Any]] = {}
    for source_row in range(7, ws.max_row + 1):
        parameter = clean_text(ws.cell(source_row, 2).value)
        if not parameter:
            continue
        operation = {
            "accelerating_voltage": clean_text(ws.cell(source_row, 14).value),
            "beam_current": clean_text(ws.cell(source_row, 15).value),
            "beam_diameter": clean_text(ws.cell(source_row, 16).value),
            "count_time": clean_text(ws.cell(source_row, 17).value),
        }
        output[parameter] = {
            "method_code": clean_text(ws.cell(source_row, 1).value),
            "detection_limit": ws.cell(source_row, 4).value,
            "detection_limit_unit": clean_text(ws.cell(source_row, 5).value),
            "procedural_blank": ws.cell(source_row, 7).value,
            "procedural_blank_unit": clean_text(ws.cell(source_row, 8).value),
            "operation": {key: value for key, value in operation.items() if value},
        }
    return output


def load_grid_context(cells: Iterable[str]) -> Dict[str, Dict[str, str]]:
    wanted = set(cells)
    columns = {
        "h3_r6",
        "state_or_ut_2011",
        "district_2011",
        "census_state_code_2011",
        "census_district_code_2011",
        "geology_age",
        "geology_supergroup",
        "geology_group",
        "geology_stratigraphy",
        "top_materials_json",
        "top_material_scores_json",
    }
    with GRID.open(newline="", encoding="utf-8") as handle:
        header = next(csv.reader(handle))
    index = {name: header.index(name) for name in columns}
    output: Dict[str, Dict[str, str]] = {}
    ripgrep = shutil.which("rg")
    if ripgrep:
        # The nationwide grid is ~470 MB and only two exact H3 rows are needed.
        # Fixed-string scanning in ripgrep avoids materializing ~16M unused fields.
        for cell in sorted(wanted):
            match = subprocess.run(
                [ripgrep, "--no-filename", "--max-count", "1", "--fixed-strings", cell, str(GRID)],
                check=True,
                capture_output=True,
                text=True,
            ).stdout
            parsed = next(csv.reader([match]))
            if parsed[index["h3_r6"]] != cell:
                raise ValueError(f"Unexpected grid match while resolving {cell}")
            output[cell] = {key: parsed[position] for key, position in index.items()}
    else:
        with GRID.open(newline="", encoding="utf-8") as handle:
            reader = csv.reader(handle)
            next(reader)
            for parsed in reader:
                cell = parsed[index["h3_r6"]]
                if cell in wanted:
                    output[cell] = {key: parsed[position] for key, position in index.items()}
                    if len(output) == len(wanted):
                        break
    missing = wanted - output.keys()
    if missing:
        raise ValueError(f"Sample H3 cells absent from KHANAN grid: {sorted(missing)}")
    return output


def classify_value(value: Any) -> Tuple[str, str, str, str]:
    source_text = numeric_text(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return source_text, source_text, "measured_numeric", ""
    lowered = source_text.lower().replace(" ", "")
    if lowered in {"b.d.l.", "bdl", "b.d.l"}:
        return source_text, "", "below_detection", "source_bdl_token_detection_limit_unavailable"
    if source_text.upper() == "NA":
        return source_text, "", "source_na_token_unresolved", "source_note_says_na_stands_for_analyzed_semantics_unresolved"
    if source_text in {"-", "–", "—"}:
        return source_text, "", "source_dash", "source_dash_not_quantified"
    if not source_text:
        return "", "", "missing", "source_blank"
    return source_text, "", "unparsed_text", "source_text_not_numeric"


def observation_rows(
    workbook,
    source_workbook: str,
    scope: str,
    samples: Mapping[str, Mapping[str, Any]],
    grid_context: Mapping[str, Mapping[str, str]],
    symbol_to_id: Mapping[str, str],
    element_names: Mapping[str, str],
    ontology_ids: set,
) -> List[Dict[str, Any]]:
    ws = workbook["3 Data"]
    parameter_start = 9 if scope == "bulk" else 15
    metadata_count = parameter_start - 1
    primary_sheet = "4 Primary Analytical Metadata" if scope == "bulk" else "4 Primary Method Metadata"
    primary = parse_primary_metadata(workbook, primary_sheet)
    method_specific = parse_method_specific(workbook)
    parameters = []
    for col in range(parameter_start, ws.max_column + 1):
        parameter = clean_text(ws.cell(2, col).value)
        if not parameter:
            continue
        parameters.append(
            {
                "col": col,
                "parameter": parameter,
                "method_code": clean_text(ws.cell(3, col).value),
                "unit": clean_text(ws.cell(4, col).value),
            }
        )
    output: List[Dict[str, Any]] = []
    analysis_sequence = 0
    for source_row in range(7, ws.max_row + 1):
        sample_name = clean_text(ws.cell(source_row, 1).value)
        if not sample_name:
            continue
        if sample_name not in samples:
            if sample_name.startswith("NA stands for") or sample_name.startswith("b.d.l."):
                continue
            raise ValueError(f"Unknown sample in {source_workbook} row {source_row}: {sample_name}")
        analysis_sequence += 1
        sample = samples[sample_name]
        cell = h3.latlng_to_cell(sample["latitude"], sample["longitude"], 6)
        context = grid_context[cell]
        mineral_code = "" if scope == "bulk" else clean_text(ws.cell(source_row, 3).value)
        mineral_name, mineral_id = MINERAL_CODE_MAP.get(mineral_code, ("", ""))
        if mineral_id and mineral_id not in ontology_ids:
            raise ValueError(f"Curated mineral mapping missing from ontology: {mineral_id}")
        source_analyzed_material = clean_text(ws.cell(source_row, 3).value) if scope == "bulk" else "mineral"
        preparation = clean_text(ws.cell(source_row, 6).value) if scope == "bulk" else clean_text(ws.cell(source_row, 4).value)
        treatment = clean_text(ws.cell(source_row, 7).value) if scope == "bulk" else clean_text(ws.cell(source_row, 5).value)
        for parameter_position, pinfo in enumerate(parameters, 1):
            parameter = pinfo["parameter"]
            value = ws.cell(source_row, pinfo["col"]).value
            source_text, numeric_value, status, qualifier = classify_value(value)
            category, chemical_name, name_basis, component_ids, parameter_ids = parameter_semantics(
                parameter, symbol_to_id, element_names
            )
            pmeta = primary.get(parameter, {})
            smeta = method_specific.get(parameter, {})
            metadata_code = clean_text(pmeta.get("metadata_method_code", smeta.get("method_code", "")))
            flags = []
            if parameter == "143Nd/144Nd" and pinfo["method_code"] != metadata_code:
                flags.append("source_method_code_conflict")
            if status == "source_na_token_unresolved":
                flags.append("source_na_token_unresolved")
            if status == "below_detection":
                flags.append("detection_limit_not_reported")
            if scope == "in_situ" and smeta.get("operation", {}).get("beam_diameter") == "1mm":
                flags.append("source_beam_diameter_1mm_retained_verbatim")
            unit_source = pinfo["unit"]
            unit_normalized = "wt.%" if unit_source in {"wt%", "wt.%"} else unit_source
            observation_id = (
                f"ECL4498-{scope.upper()}-{sample_name}-{analysis_sequence:03d}-"
                f"{parameter_position:02d}-{re.sub(r'[^A-Za-z0-9]+', '', parameter).upper()}"
            )
            row = {
                "observation_id": observation_id,
                "sample_record_id": f"ECL4498-SAMPLE-{sample_name}",
                "sample_name": sample_name,
                "source_site_id": "ECL4498-SITE-" + hashlib.sha1(
                    f"{sample['latitude']:.6f},{sample['longitude']:.6f}".encode("utf-8")
                ).hexdigest()[:10].upper(),
                "latitude": numeric_text(sample["latitude"]),
                "longitude": numeric_text(sample["longitude"]),
                "coordinate_reference_system": "EPSG:4326",
                "h3_r6": cell,
                "state_or_ut_2011": context["state_or_ut_2011"],
                "district_2011": context["district_2011"],
                "analysis_scope": scope,
                "analysis_sequence": analysis_sequence,
                "source_analyzed_material": source_analyzed_material,
                "source_mineral_code": mineral_code,
                "mineral_name_curated": mineral_name,
                "mineral_ontology_material_id": mineral_id,
                "parameter": parameter,
                "parameter_class": category,
                "chemical_name_curated": chemical_name,
                "chemical_name_basis": name_basis,
                "component_element_material_ids_json": compact_json(component_ids),
                "ontology_material_ids_json": compact_json(parameter_ids),
                "source_value_text": source_text,
                "numeric_value": numeric_value,
                "measurement_status": status,
                "measurement_qualifier": qualifier,
                "unit_source": unit_source,
                "unit_normalized": unit_normalized,
                "source_method_code_data_sheet": pinfo["method_code"],
                "source_method_code_metadata_sheet": metadata_code,
                "technique": clean_text(pmeta.get("technique", "")),
                "instrument": clean_text(pmeta.get("instrument", "")),
                "laboratory": clean_text(pmeta.get("laboratory", "")),
                "analytical_accuracy_source": compact_json(pmeta.get("accuracy", [])),
                "reference_materials_json": compact_json(pmeta.get("references", [])),
                "detection_limit_reported": str(smeta.get("detection_limit") is not None).lower(),
                "detection_limit_value": numeric_text(smeta.get("detection_limit")),
                "detection_limit_unit": clean_text(smeta.get("detection_limit_unit", "")),
                "procedural_blank_value": numeric_text(smeta.get("procedural_blank")),
                "procedural_blank_unit": clean_text(smeta.get("procedural_blank_unit", "")),
                "operation_conditions_json": compact_json(smeta.get("operation", {})),
                "sample_preparation_source": preparation,
                "chemical_treatment_source": treatment,
                "source_workbook": source_workbook,
                "source_sheet": "3 Data",
                "source_row": source_row,
                "source_column": get_column_letter(pinfo["col"]),
                "source_dataset_id": DATASET_ID,
                "source_doi": DOI,
                "source_url": DATASET_URL,
                "source_license": LICENSE,
                "source_file_sha1": SOURCE_SHA1[source_workbook],
                "source_file_sha256": SOURCE_SHA256[source_workbook],
                "source_citation": SOURCE_CITATION,
                "source_accessed_date": ACCESSED_DATE,
                "deposit_or_mine_status_verified": "false",
                "model_evidence_role": "independent_geochemical_context",
                "model_use_status": "excluded_from_v0.6_model_pending_spatial_validation",
                "data_quality_flags_json": compact_json(flags),
            }
            output.append(row)
    expected_parameter_count = 51 if scope == "bulk" else 12
    expected_analysis_count = 13 if scope == "bulk" else 55
    if len(parameters) != expected_parameter_count or analysis_sequence != expected_analysis_count:
        raise ValueError(
            f"Unexpected {scope} source shape: {len(parameters)} parameters, {analysis_sequence} analysis rows"
        )
    if len(output) != expected_parameter_count * expected_analysis_count:
        raise ValueError(f"Unexpected {scope} observation count: {len(output)}")
    return output


def write_csv(path: Path, fieldnames: Sequence[str], rows: Sequence[Mapping[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="raise", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def build_sample_rows(
    samples: Mapping[str, Mapping[str, Any]], observations: Sequence[Mapping[str, Any]], grid_context: Mapping[str, Mapping[str, str]]
) -> List[Dict[str, Any]]:
    grouped: Dict[str, List[Mapping[str, Any]]] = defaultdict(list)
    for observation in observations:
        grouped[observation["sample_name"]].append(observation)
    rows: List[Dict[str, Any]] = []
    for sample_name, sample in samples.items():
        sample_observations = grouped[sample_name]
        cell = h3.latlng_to_cell(sample["latitude"], sample["longitude"], 6)
        context = grid_context[cell]
        status_counts = Counter(item["measurement_status"] for item in sample_observations)
        scopes = defaultdict(list)
        for item in sample_observations:
            scopes[item["analysis_scope"]].append(item)
        bulk_analysis_sequences = {item["analysis_sequence"] for item in scopes["bulk"]}
        in_situ_sequences = {item["analysis_sequence"] for item in scopes["in_situ"]}
        bulk_parameters = list(dict.fromkeys(item["parameter"] for item in scopes["bulk"]))
        in_situ_parameters = list(dict.fromkeys(item["parameter"] for item in scopes["in_situ"]))
        mineral_codes = list(dict.fromkeys(item["source_mineral_code"] for item in scopes["in_situ"] if item["source_mineral_code"]))
        mineral_names = [MINERAL_CODE_MAP[code][0] for code in mineral_codes if code in MINERAL_CODE_MAP]
        ontology_materials = sorted(
            {
                material_id
                for item in sample_observations
                for material_id in json.loads(item["ontology_material_ids_json"])
                if material_id
            }
            | {MINERAL_CODE_MAP[code][1] for code in mineral_codes if code in MINERAL_CODE_MAP and MINERAL_CODE_MAP[code][1]}
        )
        chemical_names = list(dict.fromkeys(item["chemical_name_curated"] for item in sample_observations))
        key_values = {parameter: "" for parameter in KEY_TRACE_PARAMETERS}
        for item in scopes["bulk"]:
            if item["parameter"] in key_values and item["measurement_status"] == "measured_numeric":
                key_values[item["parameter"]] = item["numeric_value"]
        flags = ["targeted_petrological_sampling_not_population_representative"]
        if not sample["lithology"]:
            flags.append("sample_lithology_not_reported_in_sample_table")
        if not sample["igsn"]:
            flags.append("igsn_not_reported")
        if any(item["measurement_status"] == "source_na_token_unresolved" for item in sample_observations):
            flags.append("source_na_token_unresolved")
        source_site_id = "ECL4498-SITE-" + hashlib.sha1(
            f"{sample['latitude']:.6f},{sample['longitude']:.6f}".encode("utf-8")
        ).hexdigest()[:10].upper()
        row = {
            "record_id": f"ECL4498-SAMPLE-{sample_name}",
            "sample_name": sample_name,
            "record_class": "published_geochemical_sample",
            "latitude": numeric_text(sample["latitude"]),
            "longitude": numeric_text(sample["longitude"]),
            "coordinate_reference_system": "EPSG:4326",
            "coordinate_precision_note": "Source reports decimal-degree GPS location; collection accuracy and datum are not stated. EPSG:4326 is a KHANAN interoperability assignment.",
            "elevation_m_source": numeric_text(sample["elevation"]),
            "source_location_keyword": sample["location_keyword"],
            "source_igsn": sample["igsn"],
            "source_lithology": sample["lithology"],
            "source_site_id": source_site_id,
            "h3_r6": cell,
            "state_or_ut_2011": context["state_or_ut_2011"],
            "district_2011": context["district_2011"],
            "census_state_code_2011": context["census_state_code_2011"],
            "census_district_code_2011": context["census_district_code_2011"],
            "grid_geology_age": context["geology_age"],
            "grid_geology_supergroup": context["geology_supergroup"],
            "grid_geology_group": context["geology_group"],
            "grid_geology_stratigraphy": context["geology_stratigraphy"],
            "co_located_model_top_materials_json": context["top_materials_json"],
            "co_located_model_top_material_scores_json": context["top_material_scores_json"],
            "bulk_analysis_rows": len(bulk_analysis_sequences),
            "in_situ_analysis_spots": len(in_situ_sequences),
            "observation_rows_total": len(sample_observations),
            "numeric_observations": status_counts["measured_numeric"],
            "below_detection_observations": status_counts["below_detection"],
            "unresolved_source_na_observations": status_counts["source_na_token_unresolved"],
            "source_dash_observations": status_counts["source_dash"],
            "missing_observations": status_counts["missing"],
            "bulk_parameters_json": compact_json(bulk_parameters),
            "in_situ_parameters_json": compact_json(in_situ_parameters),
            "analyzed_mineral_codes_json": compact_json(mineral_codes),
            "analyzed_mineral_names_curated_json": compact_json(mineral_names),
            "ontology_material_ids_observed_json": compact_json(ontology_materials),
            "chemical_names_observed_json": compact_json(chemical_names),
            **{f"{parameter.lower()}_ppm_bulk": key_values[parameter] for parameter in KEY_TRACE_PARAMETERS},
            "source_sample_classification_note": "The dataset-level abstract describes Li-rich pegmatites, barren garnet-bearing pegmatites and two-mica granites; the sample table does not assign those classes to individual sample IDs.",
            "sampling_design_limitation": "Targeted petrological samples at two published coordinate sites; not a systematic regional survey and not representative of unsampled areas.",
            "deposit_or_mine_status_verified": "false",
            "model_evidence_role": "independent_geochemical_context",
            "model_use_status": "excluded_from_v0.6_model_pending_spatial_validation",
            "model_linkage_note": "H3/admin/geology/model fields are a deterministic co-location join. The source samples did not train or validate v0.6; co-located top-five scores are context, not source confirmation.",
            "source_dataset_id": DATASET_ID,
            "source_title": TITLE,
            "source_doi": DOI,
            "source_url": DATASET_URL,
            "source_authors_json": compact_json(AUTHORS),
            "source_license": LICENSE,
            "source_license_url": LICENSE_URL,
            "source_files_json": compact_json([BULK_NAME, INSITU_NAME]),
            "source_file_sha1_json": compact_json(SOURCE_SHA1),
            "source_file_sha256_json": compact_json(SOURCE_SHA256),
            "source_citation": SOURCE_CITATION,
            "source_accessed_date": ACCESSED_DATE,
            "data_quality_flags_json": compact_json(flags),
        }
        rows.append(row)
    return rows


def validate(
    samples: Sequence[Mapping[str, Any]],
    observations: Sequence[Mapping[str, Any]],
    source_shape: Mapping[str, int],
    ontology_ids: set,
) -> Dict[str, Any]:
    statuses = Counter(row["measurement_status"] for row in observations)
    scopes = Counter(row["analysis_scope"] for row in observations)
    parameter_counts = {
        scope: len({row["parameter"] for row in observations if row["analysis_scope"] == scope})
        for scope in ["bulk", "in_situ"]
    }
    unique_coordinates = {(row["latitude"], row["longitude"]) for row in samples}
    unique_ids = len({row["observation_id"] for row in observations})
    flags = Counter(
        flag for row in observations for flag in json.loads(row["data_quality_flags_json"] or "[]")
    )
    registered_source_ids = set()
    with (OUT / "source_registry.csv").open(newline="", encoding="utf-8") as handle:
        registered_source_ids = {row["source_id"] for row in csv.DictReader(handle)}
    sample_ids = {row["record_id"] for row in samples}
    observation_sample_ids = {row["sample_record_id"] for row in observations}
    linked_ontology_ids = {
        material_id
        for row in observations
        for column in ["component_element_material_ids_json", "ontology_material_ids_json"]
        for material_id in json.loads(row[column])
    } | {
        row["mineral_ontology_material_id"] for row in observations if row["mineral_ontology_material_id"]
    }
    tests = {
        "source_archive_sha1_matches_publisher": sha(RAW / ZIP_NAME, "sha1") == ZIP_SHA1,
        "source_file_sha1_matches_publisher": all(
            sha(RAW / name, "sha1") == expected for name, expected in SOURCE_SHA1.items()
        ),
        "sample_count_13": len(samples) == 13,
        "unique_coordinate_site_count_2": len(unique_coordinates) == 2,
        "bulk_analysis_rows_13": source_shape["bulk_analysis_rows"] == 13,
        "in_situ_analysis_rows_55": source_shape["in_situ_analysis_rows"] == 55,
        "bulk_parameters_51": parameter_counts["bulk"] == 51,
        "in_situ_parameters_12": parameter_counts["in_situ"] == 12,
        "observation_rows_1323": len(observations) == 1323,
        "observation_ids_unique": unique_ids == len(observations),
        "observation_sample_foreign_keys_resolve": observation_sample_ids == sample_ids,
        "source_dataset_id_resolves_in_registry": DATASET_ID in registered_source_ids
        and all(row["source_dataset_id"] == DATASET_ID for row in list(samples) + list(observations)),
        "observation_ontology_ids_resolve": linked_ontology_ids <= ontology_ids,
        "all_coordinates_in_india_bounds": all(
            6 <= float(row["latitude"]) <= 38 and 68 <= float(row["longitude"]) <= 98 for row in samples
        ),
        "all_admin_joins_present": all(row["state_or_ut_2011"] and row["district_2011"] for row in samples),
        "all_h3_cells_valid": all(h3.is_valid_cell(row["h3_r6"]) for row in samples),
        "numeric_values_nonnegative_except_end": all(
            row["parameter"] == "eNd" or float(row["numeric_value"]) >= 0
            for row in observations
            if row["measurement_status"] == "measured_numeric"
        ),
        "wt_percent_values_between_0_and_100": all(
            0 <= float(row["numeric_value"]) <= 100
            for row in observations
            if row["measurement_status"] == "measured_numeric" and row["unit_normalized"] == "wt.%"
        ),
        "measurement_status_counts_match_source": statuses
        == {
            "measured_numeric": 1254,
            "below_detection": 21,
            "source_na_token_unresolved": 2,
            "source_dash": 42,
            "missing": 4,
        },
        "isotope_ratio_method_conflict_count_13": flags["source_method_code_conflict"] == 13,
        "source_na_semantics_flag_count_2": flags["source_na_token_unresolved"] == 2,
        "raw_measurements_excluded_from_v0_6": all(
            row["model_use_status"] == "excluded_from_v0.6_model_pending_spatial_validation" for row in observations
        ),
        "no_rows_claim_verified_deposit_or_mine": all(row["deposit_or_mine_status_verified"] == "false" for row in observations),
    }
    if not all(tests.values()):
        failed = [name for name, passed in tests.items() if not passed]
        raise ValueError(f"EarthChem validation failed: {failed}")
    return {
        "dataset_version": "earthchem_ecl_4498_alpha17_v1",
        "generated_utc": "2026-09-10T00:00:00Z",
        "source_dataset_id": DATASET_ID,
        "source_doi": DOI,
        "source_license": LICENSE,
        "source_checksums": {
            ZIP_NAME: {"sha1": ZIP_SHA1, "sha256": sha(RAW / ZIP_NAME, "sha256")},
            **{
                name: {"sha1": SOURCE_SHA1[name], "sha256": SOURCE_SHA256[name]}
                for name in [BULK_NAME, INSITU_NAME]
            },
        },
        "counts": {
            "samples": len(samples),
            "unique_coordinate_sites": len(unique_coordinates),
            "source_bulk_analysis_rows": source_shape["bulk_analysis_rows"],
            "source_in_situ_analysis_rows": source_shape["in_situ_analysis_rows"],
            "observations_total": len(observations),
            "observations_by_scope": dict(sorted(scopes.items())),
            "parameters_by_scope": parameter_counts,
            "measurement_status_counts": dict(sorted(statuses.items())),
            "data_quality_flag_counts": dict(sorted(flags.items())),
        },
        "interpretation_contract": {
            "evidence_role": "Independent, CC-BY published geochemical context at two targeted coordinate sites.",
            "not_a_discovery_claim": True,
            "not_a_mine_or_deposit_register": True,
            "not_model_training_or_validation": True,
            "missing_value_policy": "Source blanks, b.d.l., NA and dash tokens remain distinct; none are converted to zero.",
            "source_na_policy": "The source note says 'NA stands for analyzed', which does not define a numeric value; numeric_value is blank and the token is flagged unresolved.",
            "detection_limit_policy": "b.d.l. is retained as below_detection, but no numeric detection limits are reported for the affected measurements.",
            "method_conflict_policy": "The Data sheet assigns method 4 to 143Nd/144Nd while Primary Metadata assigns TIMS method 3; both codes are retained and flagged.",
            "beam_diameter_policy": "The in-situ method sheet's '1mm' beam diameter is retained verbatim and flagged rather than silently corrected.",
        },
        "tests": tests,
        "status": "pass",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh", action="store_true", help="Redownload and checksum-verify the EarthChem source archive")
    args = parser.parse_args()
    ensure_source(args.refresh)
    symbol_to_id, element_names, ontology_ids = load_ontology()
    with tempfile.TemporaryDirectory(prefix="khanan-earthchem-") as temp_dir:
        converted = convert_xls_sources(Path(temp_dir))
        # Normal mode is intentionally used: these workbooks are tiny, while
        # openpyxl's read-only worksheets reparse XML for repeated random cell
        # access and turn this metadata-rich extraction into a minute-long job.
        bulk_workbook = load_workbook(converted[BULK_NAME], data_only=True, read_only=False)
        insitu_workbook = load_workbook(converted[INSITU_NAME], data_only=True, read_only=False)
        bulk_samples = parse_samples(bulk_workbook)
        insitu_samples = parse_samples(insitu_workbook)
        if bulk_samples != insitu_samples:
            raise ValueError("Bulk and in-situ source sample tables differ")
        cells = {
            h3.latlng_to_cell(sample["latitude"], sample["longitude"], 6) for sample in bulk_samples.values()
        }
        grid_context = load_grid_context(cells)
        bulk_observations = observation_rows(
            bulk_workbook,
            BULK_NAME,
            "bulk",
            bulk_samples,
            grid_context,
            symbol_to_id,
            element_names,
            ontology_ids,
        )
        insitu_observations = observation_rows(
            insitu_workbook,
            INSITU_NAME,
            "in_situ",
            bulk_samples,
            grid_context,
            symbol_to_id,
            element_names,
            ontology_ids,
        )
        observations = bulk_observations + insitu_observations
        sample_rows = build_sample_rows(bulk_samples, observations, grid_context)
        validation = validate(
            sample_rows,
            observations,
            {"bulk_analysis_rows": len(bulk_observations) // 51, "in_situ_analysis_rows": len(insitu_observations) // 12},
            ontology_ids,
        )
    write_csv(OUT / "india_earthchem_geochemical_samples.csv", SAMPLE_FIELDS, sample_rows)
    write_csv(OUT / "india_earthchem_geochemical_observations.csv", OBSERVATION_FIELDS, observations)
    (OUT / "earthchem_geochemical_validation.json").write_text(
        json.dumps(validation, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        compact_json(
            {
                "samples": len(sample_rows),
                "observations": len(observations),
                "measurement_status_counts": validation["counts"]["measurement_status_counts"],
                "validation": validation["status"],
            }
        )
    )


if __name__ == "__main__":
    main()
